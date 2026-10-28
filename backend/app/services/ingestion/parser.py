"""Document parsing & clause splitting (PDF / Word / plain text) for ingestion."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from app.core.exceptions import BadRequestError

_CN_ONLY = "零〇一二三四五六七八九十百千万两"
_CN_NUM = _CN_ONLY + "0-9"
_CHAPTER_RE = re.compile(rf"^\s*(第[{_CN_NUM}]+[编章节篇])\s*[\.、:：]?\s*(.*)$")
_ARTICLE_RE = re.compile(rf"^\s*(第[{_CN_NUM}]+条)\s*[\.、:：]?\s*(.*)$")
_ITEM_RE = re.compile(rf"^\s*[（(]([{_CN_NUM}]+)[）)]\s*[、.]?\s*(.*)$")
_CN_ITEM_RE = re.compile(rf"^\s*([{_CN_ONLY}]+)[、．.]\s*(.*)$")
_NUM_RE = re.compile(r"^\s*(\d+(?:[.．]\d+)+|\d+[、)])\s*(.*)$")
_ANNEX_RE = re.compile(r"^\s*(附件\s*[一二三四五六七八九十0-9]*|附录\s*[A-Za-z一二三四五六七八九十0-9]*)\s*[：:]?\s*(.*)$")
_MAX_CLAUSE_CHARS = 2400


@dataclass
class RawClause:
    chapter: str | None
    clause_no: str
    content: str
    source_page: int | None = None
    source_locator: str | None = None
    heading_path: list[str] = field(default_factory=list)
    block_type: str = "paragraph"


def _visible(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _split_long_content(content: str, max_chars: int = _MAX_CLAUSE_CHARS) -> list[str]:
    """Split an oversized unit on paragraph/sentence boundaries without truncation."""
    text = (content or "").strip()
    if len(_visible(text)) <= max_chars:
        return [text] if text else []
    pieces = [part.strip() for part in re.split(r"(?<=[。！？；])\s*|\n+", text) if part.strip()]
    if not pieces:
        pieces = [text]
    chunks: list[str] = []
    buf = ""
    for piece in pieces:
        while len(_visible(piece)) > max_chars:
            if buf:
                chunks.append(buf.strip())
                buf = ""
            cut = max_chars
            chunks.append(piece[:cut].strip())
            piece = piece[cut:].strip()
        candidate = f"{buf}\n{piece}".strip() if buf else piece
        if buf and len(_visible(candidate)) > max_chars:
            chunks.append(buf.strip())
            buf = piece
        else:
            buf = candidate
    if buf:
        chunks.append(buf.strip())
    return [chunk for chunk in chunks if chunk]


def _append_clause(
    clauses: list[RawClause], chapter: str | None, clause_no: str, content: str,
    *, heading_path: list[str] | None = None, block_type: str = "paragraph",
) -> None:
    parts = _split_long_content(content)
    for index, part in enumerate(parts, start=1):
        suffix = f"（续{index}）" if len(parts) > 1 else ""
        clauses.append(RawClause(
            chapter=(chapter or None), clause_no=f"{clause_no}{suffix}", content=part,
            heading_path=list(heading_path or ([chapter] if chapter else [])), block_type=block_type,
        ))


def parse_text(text: str) -> list[RawClause]:
    """Split free text into clause units.

    Recognises three heading families — 第X条/第X章, （一）（二）, and
    multi-level numeric headings (1.1 / 1.1.1) — and falls back to paragraph
    splitting when a document has no explicit numbering at all (so a legacy
    .doc with only section titles never collapses into a single "全文" blob).
    """
    clauses: list[RawClause] = []
    current_chapter: str | None = None
    heading_path: list[str] = []
    buf: list[str] = []
    current_no = ""

    def flush() -> None:
        nonlocal buf, current_no
        if buf and current_no:
            _append_clause(
                clauses, current_chapter, current_no, "\n".join(buf).strip(), heading_path=heading_path
            )
        buf = []

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        chapter_match = _CHAPTER_RE.match(line)
        if chapter_match:
            flush()
            label, rest = chapter_match.group(1), chapter_match.group(2).strip()
            current_chapter = (label + (f" {rest}" if rest else ""))[:64]
            if label.endswith(("编", "篇", "章")):
                heading_path = [current_chapter]
            else:
                heading_path = [*heading_path[:1], current_chapter]
            current_no = ""
            buf = []
            continue
        article_match = _ARTICLE_RE.match(line)
        if article_match:
            flush()
            current_no = article_match.group(1)
            rest = article_match.group(2).strip()
            buf = [rest] if rest else []
            continue
        m2 = _ITEM_RE.match(line)
        if m2:
            flush()
            current_no = f"（{m2.group(1)}）"
            buf = [m2.group(2)] if m2.group(2) else []
            continue
        m_cn = _CN_ITEM_RE.match(line)
        if m_cn:
            flush()
            current_no = f"{m_cn.group(1)}、"
            buf = [m_cn.group(2)] if m_cn.group(2) else []
            continue
        m3 = _NUM_RE.match(line)
        if m3:
            flush()
            current_no = m3.group(1)
            buf = [m3.group(2)] if m3.group(2) else []
            continue
        annex = _ANNEX_RE.match(line)
        if annex:
            flush()
            current_no = annex.group(1).replace(" ", "")
            buf = [annex.group(2)] if annex.group(2) else []
            continue
        buf.append(line)
    flush()

    if clauses:
        return clauses
    # Fallback: bounded paragraph groups. Never turn an entire long document
    # into one citable unit, but also avoid treating every wrapped PDF line as
    # a separate article.
    paragraphs = [p.strip() for p in text.splitlines() if p.strip()]
    grouped: list[str] = []
    buf_text = ""
    for paragraph in paragraphs:
        candidate = f"{buf_text}\n{paragraph}".strip() if buf_text else paragraph
        if buf_text and len(_visible(candidate)) > _MAX_CLAUSE_CHARS:
            grouped.extend(_split_long_content(buf_text))
            buf_text = paragraph
        else:
            buf_text = candidate
    if buf_text:
        grouped.extend(_split_long_content(buf_text))
    for index, paragraph in enumerate(grouped or _split_long_content(text), start=1):
        clauses.append(RawClause(
            chapter=None, clause_no=f"第{index}段", content=paragraph,
            source_locator=f"paragraph:{index}", block_type="paragraph",
        ))
    return clauses


def _extract_pdf_pages(data: bytes) -> list[str]:
    """Extract existing PDF text layers; image-only pages remain empty."""
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("PDF 解析依赖 pypdf 未安装") from exc
    try:
        reader = PdfReader(BytesIO(data))
        pages: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)
        return pages
    except PdfReadError as exc:
        raise BadRequestError(f"PDF 解析失败：文件可能已损坏或被加密（{exc}）") from exc
    except Exception as exc:  # pragma: no cover - 兜底，避免未捕获异常导致 500
        raise BadRequestError(f"PDF 解析失败：{exc}") from exc


def _extract_pdf_text(data: bytes) -> str:
    return "\n".join(_extract_pdf_pages(data))


def _anchor_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def locate_pdf_clauses(
    data: bytes, clauses: list[RawClause], extracted_pages: list[str] | None = None
) -> list[RawClause]:
    """Attach a best-effort 1-based page anchor to parsed PDF clauses."""
    pages = [_anchor_text(page) for page in (extracted_pages or _extract_pdf_pages(data))]
    for clause in clauses:
        probe = _anchor_text(clause.content)[:32]
        if not probe:
            continue
        for index, page in enumerate(pages, start=1):
            if probe in page:
                clause.source_page = index
                clause.source_locator = f"page:{index}"
                break
    return clauses


def _extract_docx_text(data: bytes) -> str:
    """Pull raw text out of a .docx; convert failures into BadRequestError."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("Word 解析依赖 python-docx 未安装") from exc
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:  # pragma: no cover - 损坏文件 / 误传的 .doc 二进制格式等
        raise BadRequestError(
            "Word 文档解析失败：请确认文件未损坏；若使用旧版 .doc 格式，请另存为 .docx 或转换为 PDF 后上传"
        ) from exc
    # Iterate paragraphs and tables in document order.  ``doc.paragraphs``
    # silently drops table content, which is often where approval limits live.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style else ""
            if style.lower().startswith("heading") or style.startswith("标题"):
                text = re.sub(r"^#+\s*", "", text)
            lines.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row_index, row in enumerate(table.rows, start=1):
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if any(cells):
                    lines.append("表格行%d｜%s" % (row_index, "｜".join(cells)))
    return "\n".join(lines)


def _extract_xlsx_rows(data: bytes) -> list[RawClause]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("Excel 解析依赖 openpyxl 未安装") from exc
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise BadRequestError(f"Excel 解析失败：{exc}") from exc
    clauses: list[RawClause] = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if not values:
                continue
            content = "｜".join(values)
            clauses.append(RawClause(
                chapter=sheet.title[:64], clause_no=f"第{row_index}行", content=content,
                source_locator=f"sheet:{sheet.title};row:{row_index}",
                heading_path=[sheet.title], block_type="table_row",
            ))
    return clauses


def _extract_xls_rows(data: bytes) -> list[RawClause]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("旧版 Excel 解析依赖 xlrd 未安装") from exc
    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise BadRequestError(f"旧版 Excel 解析失败：{exc}") from exc
    clauses: list[RawClause] = []
    for sheet in workbook.sheets():
        for row_index in range(sheet.nrows):
            values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
            if values:
                clauses.append(RawClause(
                    chapter=sheet.name[:64], clause_no=f"第{row_index + 1}行", content="｜".join(values),
                    source_locator=f"sheet:{sheet.name};row:{row_index + 1}",
                    heading_path=[sheet.name], block_type="table_row",
                ))
    return clauses


_DOC_CTRL_RE = re.compile(r"[\r\n\x07\x0b\x0c]+")
_DOC_ALLOWED = "，。、；：（）【】《》“”‘’…%0123456789.-—、()[]{}·×"


def _extract_doc_text(data: bytes) -> str:
    """Extract body text from a legacy Word binary (.doc, OLE2 compound doc).

    Word 97-2003 stores body text as UTF-16LE with ``\\r`` / ``\\x07`` /
    ``\\x0b`` / ``\\x0c`` paragraph separators. We decode, split into
    paragraphs, keep CJK-readable lines only (dropping OLE-header garbage that
    the binary stream mis-decodes into rare glyphs), and re-join with newlines.
    """
    txt = data.decode("utf-16-le", errors="ignore")
    lines: list[str] = []
    for part in _DOC_CTRL_RE.split(txt):
        s = "".join(ch for ch in part if ("\u4e00" <= ch <= "\u9fff") or ch in _DOC_ALLOWED)
        s = s.strip(" .-—、()[]{}:：\t")
        if len(s) < 2:
            continue
        cjk = sum(1 for ch in s if "\u4e00" <= ch <= "\u9fff")
        if cjk / len(s) < 0.5:
            continue
        lines.append(s)
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded document (PDF / docx / doc / txt / raw).

    Raises :class:`BadRequestError` for corrupt/unsupported files so callers can
    surface a clean 400. Used by both the rule-based splitter and the LLM-based
    splitter.
    """
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf_text(data)
    if name.endswith(".docx"):
        return _extract_docx_text(data)
    if name.endswith(".doc"):
        text = _extract_doc_text(data)
        if not text.strip():
            raise BadRequestError("旧版 .doc 文档正文提取失败，请另存为 .docx 或转换为 PDF 后上传")
        return text
    if name.endswith(".xlsx"):
        return "\n".join(clause.content for clause in _extract_xlsx_rows(data))
    if name.endswith(".xls"):
        return "\n".join(clause.content for clause in _extract_xls_rows(data))
    return data.decode("utf-8", errors="ignore")


def inspect_machine_readability(filename: str, data: bytes) -> dict:
    """Reject image-only documents before storage or model processing.

    The product accepts born-digital files and scanned files that already have
    a usable OCR text layer. It deliberately does not perform OCR itself.
    """
    name = (filename or "").lower()
    extension = Path(name).suffix
    supported = {".pdf", ".docx", ".xlsx", ".xls", ".txt"}
    if extension not in supported:
        return {
            "accepted": False,
            "code": "unsupported_format",
            "message": "仅支持可提取文字的 PDF、DOCX、XLSX、XLS、TXT 文件",
        }
    if not data:
        return {"accepted": False, "code": "empty_file", "message": "文件内容为空"}
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            page_texts = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise BadRequestError(f"PDF 可读性检查失败：{exc}") from exc
        page_count = len(page_texts)
        visible_per_page = [len(re.sub(r"\s+", "", text)) for text in page_texts]
        text_pages = sum(1 for count in visible_per_page if count >= 20)
        visible_chars = sum(visible_per_page)
        required_pages = max(1, (page_count * 7 + 9) // 10) if page_count else 1
        accepted = visible_chars >= 30 and text_pages >= required_pages
        return {
            "accepted": accepted,
            "code": "machine_readable" if accepted else "requires_ocr",
            "message": (
                "文件文字层可读取"
                if accepted
                else "该PDF没有完整可提取文字层；本系统不接收扫描原件，请先完成OCR后重新上传"
            ),
            "page_count": page_count,
            "text_page_count": text_pages,
            "visible_chars": visible_chars,
        }
    text = extract_text(filename, data)
    visible_chars = len(re.sub(r"\s+", "", text or ""))
    accepted = visible_chars >= 1
    return {
        "accepted": accepted,
        "code": "machine_readable" if accepted else "requires_text_layer",
        "message": "文件文字可读取" if accepted else "文件未提取到有效文字，请转换为文字版后重新上传",
        "visible_chars": visible_chars,
    }


def parse_pdf(data: bytes) -> list[RawClause]:
    """Extract clauses from a PDF (rule-based heading split on the raw text)."""
    pages = _extract_pdf_pages(data)
    clauses = parse_text("\n".join(pages))
    return locate_pdf_clauses(data, clauses, pages)


def parse_docx(data: bytes) -> list[RawClause]:
    """Extract clauses from a .docx (rule-based heading split on the raw text)."""
    return parse_text(_extract_docx_text(data))


def parse_doc(data: bytes) -> list[RawClause]:
    """Extract clauses from a legacy .doc (OLE2) via UTF-16LE body extraction."""
    text = _extract_doc_text(data)
    if not text.strip():
        raise BadRequestError("旧版 .doc 文档正文提取失败，请另存为 .docx 或转换为 PDF 后上传")
    return parse_text(text)


def parse_document(filename: str, data: bytes) -> list[RawClause]:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        clauses = parse_docx(data)
        for index, clause in enumerate(clauses, start=1):
            clause.source_locator = f"clause:{index}"
        return clauses
    if name.endswith(".doc"):
        clauses = parse_doc(data)
        for index, clause in enumerate(clauses, start=1):
            clause.source_locator = f"clause:{index}"
        return clauses
    if name.endswith(".xlsx"):
        return _extract_xlsx_rows(data)
    if name.endswith(".xls"):
        return _extract_xls_rows(data)
    if name.endswith(".txt"):
        clauses = parse_text(data.decode("utf-8", errors="ignore"))
        for index, clause in enumerate(clauses, start=1):
            clause.source_locator = f"clause:{index}"
        return clauses
    if data.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
        raise BadRequestError("文件内容是二进制 Office 文档，请使用正确的 .docx/.xlsx/.xls 扩展名")
    # Default: try to decode as text.
    clauses = parse_text(data.decode("utf-8", errors="ignore"))
    for index, clause in enumerate(clauses, start=1):
        clause.source_locator = f"clause:{index}"
    return clauses
