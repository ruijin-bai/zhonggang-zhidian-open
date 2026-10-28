"""Source-faithful document AST adapters for governance assets.

The AST is deliberately independent from ORM tables. Parsers preserve document
order and typed structures first; persistence, retrieval chunking and semantic
extraction happen in later stages and can therefore be rebuilt safely.
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from xml.etree import ElementTree
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Iterable

from app.core.exceptions import BadRequestError
from app.services.assets.taxonomy import ContentNodeType, ParseDisposition
from app.services.ingestion.parser import parse_document


_CHAPTER_RE = re.compile(r"^\s*(第[零〇一二两三四五六七八九十百千万\d]+[编篇章])\s*(.*)$")
_SECTION_RE = re.compile(r"^\s*(第[零〇一二两三四五六七八九十百千万\d]+节)\s*(.*)$")
_ARTICLE_RE = re.compile(r"^\s*(第[零〇一二两三四五六七八九十百千万\d]+条)\s*(.*)$")
_ITEM_RE = re.compile(r"^\s*([（(]?[零〇一二两三四五六七八九十百千万\d]+[）)、.])\s*(.*)$")
_METADATA_ONLY_EXTENSIONS = {".ai", ".zip", ".rar", ".7z"}


@dataclass
class AstNode:
    node_type: str
    label: str = ""
    title: str = ""
    text: str = ""
    source_page: int | None = None
    source_locator: str = ""
    source_bbox: dict = field(default_factory=dict)
    sheet_name: str | None = None
    cell_range: str | None = None
    structure: dict = field(default_factory=dict)
    citable: bool = True
    children: list["AstNode"] = field(default_factory=list)

    @property
    def content_hash(self) -> str:
        payload = f"{self.node_type}\n{self.label}\n{self.title}\n{self.text}".encode("utf-8")
        return hashlib.sha256(payload).hexdigest()


@dataclass
class DocumentAst:
    filename: str
    parser_name: str
    parser_version: str
    disposition: str
    root: AstNode
    source_chars: int
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class FlatAstNode:
    node: AstNode
    order_path: str
    parent_path: str | None
    ordinal: int


def _visible(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _paragraph_kind(text: str, style: str = "") -> tuple[str, str, str]:
    for regex, kind in (
        (_CHAPTER_RE, ContentNodeType.CHAPTER.value),
        (_SECTION_RE, ContentNodeType.SECTION.value),
        (_ARTICLE_RE, ContentNodeType.ARTICLE.value),
        (_ITEM_RE, ContentNodeType.ITEM.value),
    ):
        matched = regex.match(text)
        if matched:
            return kind, matched.group(1), matched.group(2).strip()
    normalized_style = style.lower()
    if normalized_style.startswith("heading") or style.startswith("标题"):
        return ContentNodeType.SECTION.value, "", text.strip()
    return ContentNodeType.PARAGRAPH.value, "", ""


def _append_hierarchical(
    root: AstNode,
    stacks: dict[str, AstNode],
    node: AstNode,
) -> None:
    if node.node_type == ContentNodeType.CHAPTER.value:
        root.children.append(node)
        stacks[ContentNodeType.CHAPTER.value] = node
        stacks.pop(ContentNodeType.SECTION.value, None)
        stacks.pop(ContentNodeType.ARTICLE.value, None)
        return
    if node.node_type == ContentNodeType.SECTION.value:
        parent = stacks.get(ContentNodeType.CHAPTER.value, root)
        parent.children.append(node)
        stacks[ContentNodeType.SECTION.value] = node
        stacks.pop(ContentNodeType.ARTICLE.value, None)
        return
    if node.node_type == ContentNodeType.ARTICLE.value:
        parent = stacks.get(ContentNodeType.SECTION.value) or stacks.get(ContentNodeType.CHAPTER.value) or root
        parent.children.append(node)
        stacks[ContentNodeType.ARTICLE.value] = node
        return
    if node.node_type == ContentNodeType.ITEM.value:
        parent = stacks.get(ContentNodeType.ARTICLE.value) or stacks.get(ContentNodeType.SECTION.value) or root
        parent.children.append(node)
        return
    parent = stacks.get(ContentNodeType.ARTICLE.value) or stacks.get(ContentNodeType.SECTION.value) or stacks.get(ContentNodeType.CHAPTER.value) or root
    parent.children.append(node)


def _parse_docx(filename: str, data: bytes) -> DocumentAst:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("Word解析依赖python-docx未安装") from exc
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise BadRequestError(f"Word文档解析失败：{exc}") from exc

    root = AstNode(
        node_type=ContentNodeType.DOCUMENT.value,
        title=Path(filename).stem,
        source_locator="word:/document/body",
        citable=False,
    )
    stacks: dict[str, AstNode] = {}
    paragraph_index = table_index = 0
    source_chars = 0
    warnings: list[str] = []

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            source_chars += len(_visible(text))
            style = (paragraph.style.name or "") if paragraph.style else ""
            kind, label, remainder = _paragraph_kind(text, style)
            is_heading = kind in {ContentNodeType.CHAPTER.value, ContentNodeType.SECTION.value}
            title = remainder if is_heading else ""
            body = "" if is_heading else (
                text[len(label):].strip() if label and text.startswith(label) else text
            )
            node = AstNode(
                node_type=kind,
                label=label,
                title=title,
                text=body,
                source_locator=f"word:/document/body/p[{paragraph_index}]",
                structure={"style": style},
                citable=not is_heading,
            )
            _append_hierarchical(root, stacks, node)
            continue
        if child.tag.endswith("}tbl"):
            table_index += 1
            table = Table(child, document)
            table_node = AstNode(
                node_type=ContentNodeType.TABLE.value,
                label=f"表{table_index}",
                source_locator=f"word:/document/body/tbl[{table_index}]",
                structure={"rows": len(table.rows), "columns": max((len(row.cells) for row in table.rows), default=0)},
                citable=False,
            )
            for row_index, row in enumerate(table.rows, start=1):
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if not any(cells):
                    continue
                row_text = "｜".join(cells)
                source_chars += len(_visible(row_text))
                row_node = AstNode(
                    node_type=ContentNodeType.TABLE_ROW.value,
                    label=f"第{row_index}行",
                    text=row_text,
                    source_locator=f"word:/document/body/tbl[{table_index}]/tr[{row_index}]",
                    structure={"values": cells, "row_index": row_index},
                )
                for column_index, value in enumerate(cells, start=1):
                    row_node.children.append(AstNode(
                        node_type=ContentNodeType.TABLE_CELL.value,
                        text=value,
                        source_locator=(
                            f"word:/document/body/tbl[{table_index}]/tr[{row_index}]/tc[{column_index}]"
                        ),
                        cell_range=f"R{row_index}C{column_index}",
                        structure={"row": row_index, "column": column_index},
                        citable=False,
                    ))
                table_node.children.append(row_node)
            _append_hierarchical(root, stacks, table_node)
    if not root.children:
        warnings.append("Word文档未提取出正文节点")
    return DocumentAst(filename, "docx_ooxml_ast", "1", ParseDisposition.STRUCTURED.value, root, source_chars, warnings)


def _parse_xlsx(filename: str, data: bytes) -> DocumentAst:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise BadRequestError("Excel解析依赖openpyxl未安装") from exc
    try:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False)
    except Exception as exc:
        raise BadRequestError(f"Excel文档解析失败：{exc}") from exc
    root = AstNode(ContentNodeType.DOCUMENT.value, title=Path(filename).stem, source_locator="excel:/", citable=False)
    source_chars = 0
    for sheet in workbook.worksheets:
        merged_ranges = [str(item) for item in sheet.merged_cells.ranges]
        sheet_node = AstNode(
            ContentNodeType.SHEET.value,
            title=sheet.title,
            source_locator=f"excel:/sheet/{sheet.title}",
            sheet_name=sheet.title,
            structure={"max_row": sheet.max_row, "max_column": sheet.max_column, "merged_ranges": merged_ranges},
            citable=False,
        )
        for row_index, row in enumerate(sheet.iter_rows(), start=1):
            nonempty = [cell for cell in row if cell.value is not None and str(cell.value).strip()]
            if not nonempty:
                continue
            values = [str(cell.value).strip() if cell.value is not None else "" for cell in row]
            start_cell, end_cell = row[0].coordinate, row[-1].coordinate
            row_text = "｜".join(values)
            source_chars += len(_visible(row_text))
            row_node = AstNode(
                ContentNodeType.TABLE_ROW.value,
                label=f"第{row_index}行",
                text=row_text,
                source_locator=f"excel:/sheet/{sheet.title}/row/{row_index}",
                sheet_name=sheet.title,
                cell_range=f"{start_cell}:{end_cell}",
                structure={"values": values, "row_index": row_index},
            )
            for cell in nonempty:
                row_node.children.append(AstNode(
                    ContentNodeType.TABLE_CELL.value,
                    text=str(cell.value).strip(),
                    source_locator=f"excel:/sheet/{sheet.title}/cell/{cell.coordinate}",
                    sheet_name=sheet.title,
                    cell_range=cell.coordinate,
                    structure={"data_type": cell.data_type, "number_format": cell.number_format},
                    citable=False,
                ))
            sheet_node.children.append(row_node)
        root.children.append(sheet_node)
    return DocumentAst(filename, "xlsx_ast", "1", ParseDisposition.STRUCTURED.value, root, source_chars)


def _parse_legacy(filename: str, data: bytes) -> DocumentAst:
    clauses = parse_document(filename, data)
    root = AstNode(ContentNodeType.DOCUMENT.value, title=Path(filename).stem, source_locator="legacy:/", citable=False)
    source_chars = 0
    stacks: dict[str, AstNode] = {}
    for index, clause in enumerate(clauses, start=1):
        source_chars += len(_visible(clause.content))
        kind, _, _ = _paragraph_kind(clause.clause_no)
        if clause.block_type == ContentNodeType.TABLE_ROW.value:
            kind = ContentNodeType.TABLE_ROW.value
        elif kind not in {ContentNodeType.ARTICLE.value, ContentNodeType.ITEM.value}:
            kind = ContentNodeType.PARAGRAPH.value
        node = AstNode(
            node_type=kind,
            label=clause.clause_no,
            text=clause.content,
            source_page=clause.source_page,
            source_locator=clause.source_locator or f"legacy:/unit/{index}",
            sheet_name=clause.chapter if clause.block_type == "table_row" else None,
            structure={"heading_path": clause.heading_path, "legacy_chapter": clause.chapter},
        )
        _append_hierarchical(root, stacks, node)
    warnings = [] if clauses else ["未提取出结构节点"]
    return DocumentAst(filename, "legacy_structured_adapter", "1", ParseDisposition.STRUCTURED.value, root, source_chars, warnings)


def _numeric_path_key(value: str) -> tuple:
    parts = re.split(r"(\d+)", value)
    return tuple(int(part) if part.isdigit() else part.lower() for part in parts)


def _parse_pptx(filename: str, data: bytes) -> DocumentAst:
    root = AstNode(
        ContentNodeType.DOCUMENT.value,
        title=Path(filename).stem,
        source_locator="pptx:/presentation",
        citable=False,
    )
    source_chars = 0
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            slide_names = sorted(
                (
                    name for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ),
                key=_numeric_path_key,
            )
            for slide_index, slide_name in enumerate(slide_names, start=1):
                xml_root = ElementTree.fromstring(archive.read(slide_name))
                texts = [
                    (element.text or "").strip()
                    for element in xml_root.iter()
                    if element.tag.rsplit("}", 1)[-1] == "t" and (element.text or "").strip()
                ]
                content = "\n".join(texts)
                source_chars += len(_visible(content))
                root.children.append(AstNode(
                    ContentNodeType.SLIDE.value,
                    label=f"第{slide_index}页",
                    text=content,
                    source_page=slide_index,
                    source_locator=f"pptx:/slides/slide{slide_index}",
                    structure={"text_blocks": texts},
                    citable=bool(content),
                ))
    except (zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise BadRequestError(f"PPTX结构解析失败：{exc}") from exc
    warnings = [] if any(node.citable for node in root.children) else ["演示文稿未提取出可引用文本"]
    return DocumentAst(filename, "pptx_ooxml_ast", "1", ParseDisposition.STRUCTURED.value, root, source_chars, warnings)


def _parse_ofd(filename: str, data: bytes) -> DocumentAst:
    root = AstNode(
        ContentNodeType.DOCUMENT.value,
        title=Path(filename).stem,
        source_locator="ofd:/document",
        citable=False,
    )
    source_chars = 0
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            content_names = sorted(
                (name for name in archive.namelist() if name.endswith("/Content.xml")),
                key=_numeric_path_key,
            )
            for page_index, content_name in enumerate(content_names, start=1):
                xml_root = ElementTree.fromstring(archive.read(content_name))
                texts = [
                    (element.text or "").strip()
                    for element in xml_root.iter()
                    if element.tag.rsplit("}", 1)[-1] == "TextCode" and (element.text or "").strip()
                ]
                content = "".join(texts)
                source_chars += len(_visible(content))
                root.children.append(AstNode(
                    ContentNodeType.PARAGRAPH.value,
                    label=f"第{page_index}页",
                    text=content,
                    source_page=page_index,
                    source_locator=f"ofd:/{content_name}",
                    structure={"text_blocks": texts},
                    citable=bool(content),
                ))
    except (zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise BadRequestError(f"OFD结构解析失败：{exc}") from exc
    warnings = [] if any(node.citable for node in root.children) else ["OFD未提取出可引用文本，需交融视觉解析"]
    return DocumentAst(filename, "ofd_xml_ast", "1", ParseDisposition.STRUCTURED.value, root, source_chars, warnings)


def parse_document_ast(filename: str, data: bytes) -> DocumentAst:
    extension = Path(filename or "").suffix.lower()
    if extension in _METADATA_ONLY_EXTENSIONS:
        root = AstNode(
            ContentNodeType.DOCUMENT.value,
            title=Path(filename).stem,
            source_locator="asset:/metadata",
            structure={"extension": extension, "size_bytes": len(data)},
            citable=False,
        )
        return DocumentAst(
            filename,
            "metadata_only",
            "1",
            ParseDisposition.METADATA_ONLY.value,
            root,
            0,
            ["该格式按原始素材保存，不生成文本引用单元"],
        )
    if extension == ".docx":
        return _parse_docx(filename, data)
    if extension == ".xlsx":
        return _parse_xlsx(filename, data)
    if extension == ".pptx":
        return _parse_pptx(filename, data)
    if extension == ".ofd":
        return _parse_ofd(filename, data)
    return _parse_legacy(filename, data)


def flatten_ast(root: AstNode) -> list[FlatAstNode]:
    flattened: list[FlatAstNode] = []

    def visit(node: AstNode, path: str, parent_path: str | None, ordinal: int) -> None:
        flattened.append(FlatAstNode(node=node, order_path=path, parent_path=parent_path, ordinal=ordinal))
        for index, child in enumerate(node.children, start=1):
            visit(child, f"{path}.{index:06d}", path, index)

    visit(root, "000001", None, 1)
    return flattened


def citable_nodes(nodes: Iterable[FlatAstNode]) -> list[FlatAstNode]:
    return [item for item in nodes if item.node.citable and bool(_visible(item.node.text))]
