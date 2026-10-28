"""V2 compound governance-asset model and parser regression tests."""
from __future__ import annotations

from io import BytesIO
import zipfile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from sqlalchemy import select

from app.db.models.asset import (
    AssetDocument,
    AssetVersion,
    AssetVersionDocument,
    CitationSpan,
    ContentNode,
    DocumentParseRun,
    KnowledgeAsset,
    RetrievalUnit,
)
from app.services.assets.document_ast import flatten_ast, parse_document_ast
from app.services.assets.migration import infer_asset_type, infer_document_role
from app.services.assets.parse_service import parse_asset_document
from app.services.assets.taxonomy import AssetType, ContentNodeType, DocumentRole
from app.services.ingestion.parser import inspect_machine_readability
from app.db.models.regulation_document import RegulationDocument


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("第一章 总则", level=1)
    document.add_paragraph("第一条 本办法适用于测试人员。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "等级"
    table.cell(0, 1).text = "条件"
    table.cell(1, 0).text = "一级"
    table.cell(1, 1).text = "三年以上"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "权限矩阵"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "审批权限"
    sheet.append(["金额", "审批人"])
    sheet.append([100, "部门负责人"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _zip_bytes(files: dict[str, str]) -> bytes:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return stream.getvalue()


def _blank_pdf_bytes() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(stream)
    return stream.getvalue()


def test_docx_ast_preserves_hierarchy_and_table_cells():
    ast = parse_document_ast("制度.docx", _docx_bytes())
    nodes = flatten_ast(ast.root)
    types = [item.node.node_type for item in nodes]
    assert ContentNodeType.CHAPTER.value in types
    assert ContentNodeType.ARTICLE.value in types
    assert ContentNodeType.TABLE.value in types
    assert types.count(ContentNodeType.TABLE_ROW.value) == 2
    assert types.count(ContentNodeType.TABLE_CELL.value) == 4
    assert all(item.node.source_locator for item in nodes)


def test_xlsx_ast_preserves_sheet_cells_and_merged_ranges():
    ast = parse_document_ast("权限.xlsx", _xlsx_bytes())
    nodes = flatten_ast(ast.root)
    sheet = next(item.node for item in nodes if item.node.node_type == ContentNodeType.SHEET.value)
    assert sheet.title == "权限矩阵"
    assert sheet.structure["merged_ranges"] == ["A1:B1"]
    assert any(item.node.cell_range == "A3" for item in nodes)
    assert all(
        item.node.source_locator
        for item in nodes
        if item.node.node_type in {ContentNodeType.TABLE_ROW.value, ContentNodeType.TABLE_CELL.value}
    )


def test_pptx_and_ofd_adapters_preserve_page_locators():
    pptx = _zip_bytes({
        "ppt/slides/slide1.xml": '<p:sld xmlns:p="p" xmlns:a="a"><a:t>宣贯要点</a:t></p:sld>',
        "ppt/slides/slide2.xml": '<p:sld xmlns:p="p" xmlns:a="a"><a:t>办理流程</a:t></p:sld>',
    })
    ppt_ast = parse_document_ast("宣贯.pptx", pptx)
    assert [node.source_page for node in ppt_ast.root.children] == [1, 2]
    assert ppt_ast.root.children[0].text == "宣贯要点"

    ofd = _zip_bytes({
        "Doc_0/Pages/Page_0/Content.xml": '<ofd:Page xmlns:ofd="ofd"><ofd:TextCode>印发通知</ofd:TextCode></ofd:Page>',
    })
    ofd_ast = parse_document_ast("通知.ofd", ofd)
    assert ofd_ast.root.children[0].text == "印发通知"
    assert ofd_ast.root.children[0].source_locator.endswith("Content.xml")


def test_controlled_taxonomy_inference_handles_non_regulation_assets():
    assert infer_asset_type("安全制度汇编") == AssetType.COMPILATION.value
    assert infer_asset_type("项目策划模板2025") == AssetType.TEMPLATE_PACKAGE.value
    document = RegulationDocument(
        file_id="legacy", doc_type="attachment", title="流程",
        filename="附件3-审批流程图.docx", mime="application/docx", sort_order=1,
    )
    assert infer_document_role(document) == DocumentRole.FLOWCHART.value


def test_machine_readability_rejects_pdf_without_text_layer():
    result = inspect_machine_readability("扫描制度.pdf", _blank_pdf_bytes())
    assert result["accepted"] is False
    assert result["code"] == "requires_ocr"
    assert "系统不接收扫描原件" in result["message"]


async def test_parse_snapshot_builds_nodes_citations_and_retrieval_units(session, monkeypatch):
    payload = _docx_bytes()

    class FakeStore:
        @staticmethod
        def get_object(_: str) -> bytes:
            return payload

    monkeypatch.setattr("app.services.assets.parse_service.get_object_store", lambda: FakeStore())
    asset = KnowledgeAsset(asset_type="management_method", title="测试办法")
    session.add(asset)
    await session.flush()
    version = AssetVersion(asset_id=asset.asset_id, version_label="v1")
    document = AssetDocument(
        filename="测试办法.docx", title="测试办法", mime="application/docx",
        object_key="local://test.docx", sha256="a" * 64, size_bytes=len(payload),
        parse_disposition="structured",
    )
    session.add_all([version, document])
    await session.flush()
    session.add(AssetVersionDocument(
        version_id=version.version_id, document_id=document.document_id,
        role=DocumentRole.NORMATIVE_BODY.value, mandatory=True, normative=True,
    ))
    await session.commit()

    result = await parse_asset_document(session, document.document_id)
    await session.commit()

    assert result["status"] == "succeeded"
    assert result["anchor_ratio"] == 1.0
    assert result["version_states"][version.version_id] == "ready"
    assert len((await session.execute(select(ContentNode))).scalars().all()) >= 9
    assert len((await session.execute(select(CitationSpan))).scalars().all()) == 3
    assert len((await session.execute(select(RetrievalUnit))).scalars().all()) == 3
    assert (await session.execute(select(DocumentParseRun.is_current))).scalar_one() is True

    second = await parse_asset_document(session, document.document_id)
    await session.commit()
    runs = list((await session.execute(
        select(DocumentParseRun).where(DocumentParseRun.document_id == document.document_id)
    )).scalars())
    assert second["status"] == "succeeded"
    assert sum(run.is_current for run in runs) == 1
    statuses = list((await session.execute(select(RetrievalUnit.status))).scalars())
    assert statuses.count("isolated") == 3
    assert statuses.count("verified") == 3


async def test_parse_snapshot_blocks_scanned_pdf_without_ocr(session, monkeypatch):
    payload = _blank_pdf_bytes()

    class FakeStore:
        @staticmethod
        def get_object(_: str) -> bytes:
            return payload

    monkeypatch.setattr("app.services.assets.parse_service.get_object_store", lambda: FakeStore())
    asset = KnowledgeAsset(asset_type="regulation", title="扫描制度")
    session.add(asset)
    await session.flush()
    version = AssetVersion(asset_id=asset.asset_id, version_label="v1")
    document = AssetDocument(
        filename="扫描制度.pdf", title="扫描制度", mime="application/pdf",
        object_key="local://scan.pdf", sha256="b" * 64, size_bytes=len(payload),
        parse_disposition="structured",
    )
    session.add_all([version, document])
    await session.flush()
    session.add(AssetVersionDocument(
        version_id=version.version_id, document_id=document.document_id,
        role=DocumentRole.NORMATIVE_BODY.value, mandatory=True, normative=True,
    ))
    await session.commit()

    result = await parse_asset_document(session, document.document_id)
    await session.commit()

    assert result["status"] == "requires_ocr"
    assert result["disposition"] == "excluded"
    assert result["version_states"][version.version_id] == "blocked"
    assert (await session.get(AssetDocument, document.document_id)).disposition_reason
    assert not list((await session.execute(
        select(RetrievalUnit).where(RetrievalUnit.version_id == version.version_id)
    )).scalars())
