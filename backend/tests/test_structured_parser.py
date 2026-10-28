from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.services.ingestion.parser import extract_text, parse_document, parse_text
from app.services.ingestion.quality import assess_parse_quality


def test_unnumbered_long_document_never_becomes_fulltext_clause():
    text = "\n".join([f"这是第{i}段制度正文，说明审批要求和业务证据。" * 20 for i in range(20)])
    clauses = parse_text(text)
    assert len(clauses) > 1
    assert all(item.clause_no != "全文" for item in clauses)
    assert max(len(item.content) for item in clauses) <= 2600


def test_docx_keeps_table_rows_in_document_order():
    document = Document()
    document.add_paragraph("第一章 总则")
    document.add_paragraph("第一条 本制度适用于测试。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "金额"
    table.cell(0, 1).text = "审批人"
    table.cell(1, 0).text = "100万元"
    table.cell(1, 1).text = "部门负责人"
    stream = BytesIO()
    document.save(stream)
    text = extract_text("rules.docx", stream.getvalue())
    assert "表格行1｜金额｜审批人" in text
    assert "表格行2｜100万元｜部门负责人" in text


def test_xlsx_is_parsed_as_rows_not_zip_binary():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "审批权限"
    sheet.append(["事项", "限额", "审批人"])
    sheet.append(["采购", 500000, "分管领导"])
    sheet.append(["   ", None, ""])
    stream = BytesIO()
    workbook.save(stream)
    data = stream.getvalue()
    clauses = parse_document("authority.xlsx", data)
    assert [item.block_type for item in clauses] == ["table_row", "table_row"]
    assert clauses[1].source_locator == "sheet:审批权限;row:2"
    assert "采购｜500000｜分管领导" == clauses[1].content
    assert not extract_text("authority.xlsx", data).startswith("PK")


def test_quality_gate_rejects_binary_and_giant_blob():
    clauses = parse_text("PK[Content_Types].xml" + "制度" * 5000)
    report = assess_parse_quality(
        "PK[Content_Types].xml" + "制度" * 5000,
        clauses,
        backend="native",
        filename="bad.txt",
    )
    assert report.passed is False
    assert any("二进制" in warning for warning in report.warnings)


def test_numbering_families_and_hierarchy_are_recognised():
    text = "第一章 总则\n第一条 适用范围。\n一、基本要求。\n1.1 详细要求。\n附件1：材料清单"
    clauses = parse_text(text)
    assert [item.clause_no for item in clauses] == ["第一条", "一、", "1.1", "附件1"]
    assert clauses[0].heading_path == ["第一章 总则"]
